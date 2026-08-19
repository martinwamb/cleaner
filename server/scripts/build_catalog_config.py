"""Enrich the operator workbook and generate the server's catalog snapshot.

The workbook remains the operator editing surface. Running this script after a
catalog edit refreshes the JSON consumed by the quote service.
"""

from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path
from typing import cast

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


READY_RULES = {
    "turnover": {
        "platform_id": "turnover",
        "platform_name": "Apartment turnover",
        "base": 180,
        "size_rate": 48,
        "minimum": 240,
        "size_unit": "rooms",
        "spread": 0.15,
        "condition": {"Standard": 1, "Heavy": 1.25, "Extreme": 1.5},
        "recurring": 0.9,
        "travel_type": "Zone-based",
        "travel_fee": 35,
        "basis": "Base price + rooms x unit rate, then condition, recurring, add-on, and travel rules.",
    },
    "res-deep": {
        "platform_id": "deep",
        "platform_name": "Residential deep cleaning",
        "base": 150,
        "size_rate": 42,
        "minimum": 180,
        "size_unit": "rooms",
        "spread": 0.15,
        "condition": {"Standard": 1, "Heavy": 1.25, "Extreme": 1.5},
        "recurring": 0.9,
        "travel_type": "Zone-based",
        "travel_fee": 35,
        "basis": "Base price + rooms x unit rate, then condition, recurring, add-on, and travel rules.",
    },
    "construction-final": {
        "platform_id": "construction",
        "platform_name": "Post-Construction Final Clean",
        "base": 280,
        "size_rate": 0.22,
        "minimum": 450,
        "size_unit": "sq ft",
        "spread": 0.18,
        "condition": {"Standard": 1, "Heavy": 1.2, "Extreme": 1.4},
        "recurring": 0.9,
        "travel_type": "Zone-based",
        "travel_fee": 35,
        "basis": "Base price + square feet x unit rate, then condition, recurring, add-on, and travel rules.",
    },
    "commercial-recurring": {
        "platform_id": "commercial",
        "platform_name": "Recurring Commercial Cleaning",
        "base": 220,
        "size_rate": 0.18,
        "minimum": 280,
        "size_unit": "sq ft",
        "spread": 0.18,
        "condition": {"Standard": 1, "Heavy": 1.2, "Extreme": 1.35},
        "recurring": 0.9,
        "travel_type": "Zone-based",
        "travel_fee": 35,
        "basis": "Base price + square feet x unit rate, then condition, recurring, add-on, and travel rules.",
    },
}

ADD_ON_PRICES = {
    "Inside appliances": 45,
    "Interior windows": 35,
    "Inside cabinets": 55,
}

NEW_COLUMNS = [
    "Record Type",
    "Card Icon",
    "Featured Order",
    "Pricing Model",
    "Pricing Basis",
    "Size Input Label",
    "Base Price",
    "Estimate Spread",
    "Standard Multiplier",
    "Heavy Multiplier",
    "Extreme Multiplier",
    "One-Time Multiplier",
    "Recurring Multiplier",
    "Rush Multiplier",
    "Travel Fee Type",
    "Travel Fee Amount",
    "Service Area Rule",
    "Add-on Rules",
    "Pricing Readiness",
    "Pricing Version",
    "Effective Date",
    "Last Changed By",
    "Change Reason",
    "Operator Change Guide",
    "Price Range Role",
]

VALIDATION_VALUES = {
    "Status": '"Approved,Draft,Paused,Archived,Internal only"',
    "Featured": '"TRUE,FALSE"',
    "Record Type": '"Service,Add-on,Modifier"',
    "Pricing Model": '"Flat range,Per unit,Per room,Per square foot,Hourly,Monthly contract,Custom quote"',
    "Pricing Unit": '"Per visit,Per room,Per square foot,Per unit,Hourly,Per project,Custom quote"',
    "Travel Fee Type": '"Included,Flat fee,Per mile,Zone-based,Manual review"',
    "Pricing Readiness": '"Ready,Needs operator pricing review,Blocked"',
}


def split_values(value):
    return [item.strip() for item in str(value or "").split(";") if item.strip()]


def pricing_model(unit):
    unit = str(unit or "")
    if "square foot" in unit:
        return "Per square foot"
    if "room" in unit:
        return "Per room"
    if "hour" in unit:
        return "Hourly"
    if "monthly" in unit:
        return "Monthly contract"
    if "custom" in unit.lower():
        return "Custom quote"
    return "Per unit"


def size_label(unit):
    unit = str(unit or "")
    if "square foot" in unit:
        return "Square feet"
    if "room" in unit:
        return "Rooms"
    if "hour" in unit:
        return "Hours"
    if "pane" in unit:
        return "Panes"
    if "item" in unit:
        return "Items"
    return "Units or project scope"


def add_on_rules(names, ready):
    if not ready:
        return "Needs operator pricing review"
    # The first quote release exposes these same three add-ons for each enabled
    # service. Keep the workbook and the quote form aligned until add-ons gain
    # service-specific availability in a later catalog revision.
    names = "Inside appliances; Interior windows; Inside cabinets"
    rules = []
    for name in split_values(names):
        if name in ADD_ON_PRICES:
            rules.append(f"{name} | flat | {ADD_ON_PRICES[name]}")
    return "; ".join(rules) or "None configured"


def parse_add_on_rules(value):
    rules = []
    for raw_rule in split_values(str(value or "").replace("; ", ";")):
        parts = [part.strip() for part in raw_rule.split("|")]
        if len(parts) != 3:
            continue
        try:
            price = float(parts[2])
        except ValueError:
            continue
        rules.append({"name": parts[0], "valueType": parts[1], "price": price})
    return rules


def header_map(ws):
    return {str(cell.value): cast(int, cell.column) for cell in ws[1] if cell.value is not None}


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_catalog_config.py <workbook.xlsx> <catalog-config.json>")

    workbook_path = Path(sys.argv[1]).resolve()
    config_path = Path(sys.argv[2]).resolve()
    wb = load_workbook(workbook_path)
    ws = wb["Service Catalog"]

    columns = header_map(ws)

    # The original workbook was missing a value for Travel Fee Applicable, so
    # the final five customer-facing fields were shifted one column left.
    # Repair that once before enriching the table so the generated config has
    # the intended semantics.
    for row in range(2, ws.max_row + 1):
        old_travel = ws.cell(row, columns["Travel Fee Applicable"]).value
        old_add_ons = ws.cell(row, columns["Add-ons"]).value
        old_repeat = ws.cell(row, columns["Repeat Potential"]).value
        old_note = ws.cell(row, columns["Customer-Facing Estimate Note"]).value
        old_featured = ws.cell(row, columns["Featured"]).value
        if isinstance(old_travel, str) and ";" in old_travel and old_add_ons in {"Low", "Medium", "High"}:
            ws.cell(row, columns["Travel Fee Applicable"], "Yes" if ws.cell(row, columns["Service ID"]).value in READY_RULES else "Manual review")
            ws.cell(row, columns["Add-ons"], old_travel)
            ws.cell(row, columns["Repeat Potential"], old_add_ons)
            ws.cell(row, columns["Customer-Facing Estimate Note"], old_repeat)
            ws.cell(row, columns["Featured"], old_note)

    for name in NEW_COLUMNS:
        if name not in columns:
            ws.cell(1, ws.max_column + 1, name)
    columns = header_map(ws)

    for name, formula in VALIDATION_VALUES.items():
        if name not in columns:
            continue
        for validation in list(ws.data_validations.dataValidation):
            if validation.formula1 == formula:
                ws.data_validations.dataValidation.remove(validation)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = "Choose a value from the list."
        dv.errorTitle = "Invalid catalog value"
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(2, columns[name]).coordinate}:{ws.cell(ws.max_row, columns[name]).coordinate}")

    for row in range(2, ws.max_row + 1):
        service_id = ws.cell(row, columns["Service ID"]).value
        rule = READY_RULES.get(service_id)
        unit = ws.cell(row, columns["Pricing Unit"]).value
        model = pricing_model(unit)
        existing_version = ws.cell(row, columns["Pricing Version"]).value

        def set_default(name, value):
            if ws.cell(row, columns[name]).value in (None, ""):
                ws.cell(row, columns[name], value)

        if rule and not existing_version:
            set_default("Pricing Unit", "Per room" if rule["size_unit"] == "rooms" else "Per square foot")
            set_default("Unit Rate Low", rule["size_rate"])
            set_default("Unit Rate High", rule["size_rate"])
            set_default("Minimum Price", rule["minimum"])
            set_default("Add-ons", "Inside appliances; Interior windows; Inside cabinets")
            unit = ws.cell(row, columns["Pricing Unit"]).value
            model = pricing_model(unit)
        ws.cell(row, columns["Record Type"], "Service")
        set_default("Card Icon", {"turnover": "↻", "res-deep": "✦", "construction-final": "⌂", "commercial-recurring": "▦"}.get(service_id, "▦"))
        set_default("Featured Order", {"turnover": 1, "res-deep": 2, "construction-final": 3, "commercial-recurring": 4}.get(service_id, 0))
        set_default("Pricing Model", model)
        set_default("Pricing Basis", rule["basis"] if rule else "Operator must define the base, unit, condition, frequency, add-on, and travel rules before enabling quotes.")
        set_default("Size Input Label", size_label(unit))
        set_default("Base Price", rule["base"] if rule else None)
        set_default("Estimate Spread", rule["spread"] if rule else None)
        set_default("Standard Multiplier", rule["condition"]["Standard"] if rule else None)
        set_default("Heavy Multiplier", rule["condition"]["Heavy"] if rule else None)
        set_default("Extreme Multiplier", rule["condition"]["Extreme"] if rule else None)
        set_default("One-Time Multiplier", 1 if rule else None)
        set_default("Recurring Multiplier", rule["recurring"] if rule else None)
        set_default("Rush Multiplier", None)
        set_default("Travel Fee Type", rule["travel_type"] if rule else "Manual review")
        set_default("Travel Fee Amount", rule["travel_fee"] if rule else None)
        set_default("Service Area Rule", "Twin Cities local postal-code allowlist; outside area adds the travel fee." if rule else "Define Twin Cities eligibility and out-of-area response.")
        set_default("Add-on Rules", add_on_rules(ws.cell(row, columns["Add-ons"]).value, bool(rule)))
        set_default("Pricing Readiness", "Ready" if rule else "Needs operator pricing review")
        set_default("Pricing Version", "v1.0" if rule else None)
        set_default("Effective Date", date(2026, 8, 19) if rule else None)
        set_default("Last Changed By", "Operator" if rule else None)
        set_default("Change Reason", "Aligned to quote generator rules." if rule else None)
        set_default("Operator Change Guide", "Edit Base Price, Unit Rate, Minimum Price, or a multiplier in this row. Keep Pricing Basis and Change Reason current; run the import command before testing quotes." if rule else "Complete pricing inputs and mark Ready only after scope, exclusions, and estimate behavior are approved.")
        set_default("Price Range Role", "Reference market range; quote uses the explicit calculation inputs in this row." if rule else "Research reference only; this service is not enabled for automatic quotes.")

    for cell in ws[1]:
        if cell.value in NEW_COLUMNS:
            cell.fill = PatternFill("solid", fgColor="314A46")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.comment = Comment("Operator pricing control. Changes take effect in the quote service after the catalog import command is run.", "Cleaner")
    ws.row_dimensions[1].height = 42
    for name in NEW_COLUMNS:
        ws.column_dimensions[ws.cell(1, columns[name]).column_letter].width = 23
    for row in range(2, ws.max_row + 1):
        ws.cell(row, columns["Estimate Spread"]).number_format = "0%"
        for name in ["Base Price", "Travel Fee Amount"]:
            ws.cell(row, columns[name]).number_format = '$#,##0.00;[Red]-$#,##0.00'
        for name in ["Standard Multiplier", "Heavy Multiplier", "Extreme Multiplier", "One-Time Multiplier", "Recurring Multiplier", "Rush Multiplier"]:
            ws.cell(row, columns[name]).number_format = "0.00x"
        ws.cell(row, columns["Effective Date"]).number_format = "yyyy-mm-dd"

    table = ws.tables["ServiceCatalog"]
    table.ref = f"A1:{ws.cell(ws.max_row, ws.max_column).coordinate}"
    wb.save(workbook_path)

    headers = [cell.value for cell in ws[1]]
    services = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        item = {headers[index]: value for index, value in enumerate(row)}
        rule = READY_RULES.get(item["Service ID"])
        if rule:
            item["platformId"] = rule["platform_id"]
            item["platformName"] = rule["platform_name"]
            item["sizeUnit"] = rule["size_unit"]
        item["propertyTypes"] = split_values(item["Property Types"])
        item["customerTypes"] = split_values(item["Customer Types"])
        item["frequencyOptions"] = split_values(item["Frequency Options"])
        item["quoteEnabled"] = item["Status"] == "Approved" and item["Pricing Readiness"] == "Ready"
        item["addOnRules"] = parse_add_on_rules(item["Add-on Rules"])
        services.append(item)

    configured_add_ons = {}
    for service in services:
        for add_on in service["addOnRules"]:
            configured_add_ons[add_on["name"]] = add_on
    service_by_id = {service["Service ID"]: service for service in services}

    config = {
        "generatedAt": date.today().isoformat(),
        "sourceWorkbook": workbook_path.name,
        "services": services,
        "conditions": ["Standard", "Heavy", "Extreme"],
        "frequencies": ["One-time", "Recurring"],
        "propertyTypes": sorted({value for service in services for value in service["propertyTypes"]}),
        "addOns": list(configured_add_ons.values()),
        "localPostalCodes": [
            "55401", "55402", "55403", "55404", "55405", "55406", "55407", "55408", "55409", "55410",
            "55411", "55412", "55413", "55414", "55415", "55416", "55417", "55418", "55419", "55423",
            "55424", "55425", "55426", "55429", "55430", "55435", "55436", "55439", "55450", "55454", "55455",
        ],
        "pricing": {
            service_id: {
                "base": service_by_id[service_id]["Base Price"],
                "sizeRate": service_by_id[service_id]["Unit Rate Low"],
                "minimum": service_by_id[service_id]["Minimum Price"],
                "spread": service_by_id[service_id]["Estimate Spread"],
                "condition": {
                    "Standard": service_by_id[service_id]["Standard Multiplier"],
                    "Heavy": service_by_id[service_id]["Heavy Multiplier"],
                    "Extreme": service_by_id[service_id]["Extreme Multiplier"],
                },
                "recurring": service_by_id[service_id]["Recurring Multiplier"],
                "travelFee": service_by_id[service_id]["Travel Fee Amount"],
                "addOns": service_by_id[service_id]["addOnRules"],
            }
            for service_id, rule in READY_RULES.items()
        },
    }
    config_path.parent.mkdir(parents=True, exist_ok=True)
    config_path.write_text(json.dumps(config, indent=2, default=str) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
